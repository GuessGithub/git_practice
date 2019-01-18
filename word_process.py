# -*- encoding:utf-8 -*-

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT

from docx.oxml.shared import OxmlElement, qn # make the content of cell center


doc = docx.Document('demo2.docx')
table = doc.tables

tab = table[0]

tab.cell(0,1).text='Loving you, Lucia'

# make the content of cell center

tc = tab.cell(0,1)._tc
tcPr = tc.get_or_add_tcPr()
tcVAlign = OxmlElement('w:vAlign')
tcVAlign.set(qn('w:val'), "center")
tcPr.append(tcVAlign)

for i in range(0,4):
    print(tab.row_cells(i))
    

doc.save('demo2.docx')

print(u'哈哈')
